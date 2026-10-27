"""
Экспорт данных из БД в Excel — с группировкой по кабинетам.

Формат: для каждого кабинета — отдельный блок:
  - Строка-заголовок «Каб. 102/1» (объединённая, цветная)
  - Строка с названиями колонок
  - Строки с устройствами этого кабинета
  - Пустая строка-разделитель

Колонки:
  Тип | Марка, модель | Инв. № | Сер. № | ОС | Dr.Web |
  IP-адрес | Имя ПК | Физ. MAC | Реал. MAC | Характеристики | Дополнительно | АРМ
"""
import os
import sqlite3
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config

# Колонки данных (без № и кабинета — они в заголовке блока)
HEADERS = [
    "Тип",                     # 1
    "Марка, модель",           # 2
    "Инв. №",                  # 3
    "Сер. №",                  # 4
    "ОС",                      # 5
    "Dr.Web",                  # 6
    "IP-адрес",                # 7
    "Имя ПК",                  # 8
    "Физ. MAC",                # 9
    "Реал. MAC",               # 10
    "Характеристики",          # 11
    "Дополнительно",           # 12
    "АРМ",                     # 13
]

COLUMN_MAP = [
    "tech_type",
    "tech",
    "inv_number",
    "serial",
    "os",
    "drweb",
    "ip",
    "pc_name",
    "mac_phys",
    "mac_real",
    "specs",
    "note",
    "arm_id",
]

TECH_TYPE_LABELS = {
    "pc":      "Системный блок",
    "monitor": "Монитор",
    "printer": "Принтер / МФУ",
    "laptop":  "Ноутбук",
    "other":   "Другое",
}

# === ЦВЕТА ===
CAB_HEADER_FILL = PatternFill(start_color="2F543B", end_color="2F543B", fill_type="solid")
COL_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
SERIAL_INV_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
IP_MAC_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ARM_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

CAB_HEADER_FONT = Font(bold=True, size=12, color="FFFFFF")
COL_HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
CELL_FONT = Font(size=10)

CELL_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)

BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)

COLUMN_WIDTHS = [18, 28, 20, 20, 20, 18, 16, 20, 20, 20, 28, 28, 8]

DRWEB_DISPLAY_MAP = {
    "Установлен":                "Установлен",
    "Нет лицензии":              "Нет лицензии",
    "Нет соединения с сервером": "Нет соединения с сервером",
    "Не установлен":             "Не установлен",
    "+":         "Установлен",
    "да":        "Установлен",
    "yes":       "Установлен",
    "нет лиц":   "Нет лицензии",
    "нет соед":  "Нет соединения с сервером",
    "нет":       "Не установлен",
    "no":        "Не установлен",
}


def _human_tech_type(value: str) -> str:
    if not value:
        return ""
    return TECH_TYPE_LABELS.get(value, value)


def _human_drweb(value: str) -> str:
    if not value:
        return ""
    return DRWEB_DISPLAY_MAP.get(str(value).strip(), str(value))


def _cell_value(row: dict, db_field: str) -> str:
    val = row.get(db_field) or ""
    if db_field == "mac_phys" and not val:
        val = row.get("mac") or ""
    if db_field == "tech_type":
        val = _human_tech_type(val)
    if db_field == "drweb":
        val = _human_drweb(val)
    if db_field == "arm_id" and val:
        val = f"#{val}"
    return str(val)


def export_to_excel(output_path: str = None) -> str:
    """Экспорт с группировкой по кабинетам."""
    if output_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = f"data/inventory_export_{timestamp}.xlsx"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    conn = sqlite3.connect(config.DB_PATH)
    conn.row_factory = sqlite3.Row
    cursor = conn.execute(
        "SELECT * FROM inventory ORDER BY clinic, cabinet, id"
    )
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()

    if not rows:
        print("База данных пуста!")
        return None

    wb = Workbook()
    wb.remove(wb.active)

    # Группируем по поликлиникам
    clinics = []
    seen = set()
    for r in rows:
        c = r.get("clinic") or "—"
        if c not in seen:
            seen.add(c)
            clinics.append(c)

    from config import CLINICS

    for clinic in clinics:
        clinic_rows = [r for r in rows if (r.get("clinic") or "—") == clinic]
        sheet_title = (clinic or "—")[:31]
        ws = wb.create_sheet(title=sheet_title)

        row_idx = 1
        num_cols = len(HEADERS)

        # Группируем по кабинетам
        cabinets = []
        cab_seen = set()
        for r in clinic_rows:
            cab = r.get("cabinet") or "—"
            if cab not in cab_seen:
                cab_seen.add(cab)
                cabinets.append(cab)

        for cab in cabinets:
            cab_rows = [r for r in clinic_rows if (r.get("cabinet") or "—") == cab]

            # === Строка-заголовок кабинета ===
            ws.merge_cells(start_row=row_idx, start_column=1,
                           end_row=row_idx, end_column=num_cols)
            cell = ws.cell(row=row_idx, column=1, value=f"Каб. {cab}")
            cell.fill = CAB_HEADER_FILL
            cell.font = CAB_HEADER_FONT
            cell.alignment = Alignment(horizontal='left', vertical='center')
            # Рамка для объединённой строки
            for c in range(1, num_cols + 1):
                ws.cell(row=row_idx, column=c).border = BORDER

            row_idx += 1

            # === Строка с названиями колонок ===
            for col, header in enumerate(HEADERS, 1):
                cell = ws.cell(row=row_idx, column=col, value=header)
                cell.fill = COL_HEADER_FILL
                cell.font = COL_HEADER_FONT
                cell.alignment = HEADER_ALIGN
                cell.border = BORDER

            row_idx += 1

            # === Строки с устройствами ===
            for row in cab_rows:
                for col_idx, db_field in enumerate(COLUMN_MAP, 1):
                    val = _cell_value(row, db_field)
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.alignment = CELL_ALIGN
                    cell.border = BORDER
                    cell.font = CELL_FONT

                    # Подсветка
                    if db_field in ("serial", "inv_number"):
                        cell.fill = SERIAL_INV_FILL
                    elif db_field in ("ip", "mac_phys", "mac_real") and val:
                        cell.fill = IP_MAC_FILL
                    elif db_field == "arm_id" and val:
                        cell.fill = ARM_FILL

                row_idx += 1

            # === Пустая строка-разделитель ===
            row_idx += 1

        # === Ширина колонок ===
        for i, width in enumerate(COLUMN_WIDTHS, 1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = 'A2'

    wb.save(output_path)
    print(f"Экспорт завершён: {output_path}")
    print(f"Листов: {len(clinics)}")
    return output_path


if __name__ == "__main__":
    export_to_excel()
    print("Готово!")
